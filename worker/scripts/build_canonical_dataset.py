from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from itertools import pairwise
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.datetime import from_excel  # type: ignore[import-untyped]

type JsonScalar = str | int | float | bool | None
type JsonValue = JsonScalar | Sequence[JsonValue] | Mapping[str, JsonValue]
type JsonObject = dict[str, JsonValue]
type CellValue = str | int | float | datetime | date | None

WORKER_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = WORKER_ROOT.parent
DEFAULT_SOURCE_ROOT = PROJECT_ROOT / "dataset_domrock"
DEFAULT_OUTPUT_DIR = WORKER_ROOT / "sandbox" / "data" / "domrock"

RH_FIELD_ORDER = (
    "competencia",
    "data_ref",
    "cod_marca",
    "descr_marca",
    "cod_loja",
    "descr_loja",
    "matricula",
    "data_admiss",
    "data_demiss",
    "cod_cargo",
    "descr_cargo",
)
VENDAS_FIELD_ORDER = (
    "competencia",
    "data_ref",
    "data_venda",
    "cod_marca",
    "descr_marca",
    "cod_loja",
    "descr_loja",
    "matricula",
    "vlr_venda",
)
COMISSOES_FIELD_ORDER = (
    "competencia",
    "cod_marca",
    "descr_marca",
    "cod_cargo",
    "descr_cargo",
    "percentual_comissao",
)

RH_COLUMN_MAP = {
    "Data_Ref": "data_ref",
    "Cod_Marca": "cod_marca",
    "Descri_Marca": "descr_marca",
    "Cod_Loja": "cod_loja",
    "Descr_Loja": "descr_loja",
    "Matricula": "matricula",
    "Data_Admiss": "data_admiss",
    "Data_Demiss": "data_demiss",
    "Cod_Cargo": "cod_cargo",
    "Descri_Cargo": "descr_cargo",
}
VENDAS_COLUMN_MAP = {
    "Date_Ref": "data_ref",
    "Cod_Marca": "cod_marca",
    "Descr_Marca": "descr_marca",
    "Cod_Loja": "cod_loja",
    "Descr_Loja": "descr_loja",
    "Matricula": "matricula",
    "Vlr _Venda": "vlr_venda",
}
COMISSOES_COLUMN_MAP = {
    "Cod_Marca": "cod_marca",
    "Descr_Marca": "descr_marca",
    "Cod_Cargo": "cod_cargo",
    "Descri_Cargo": "descr_cargo",
    "%_Comiss": "percentual_comissao",
}
KNOWN_CARGO_CODES = frozenset({100, 150, 200, 300})
PUBLISHED_COMPETENCIAS = ("2025-08", "2025-09", "2025-10", "2025-11", "2025-12")
RH_KEY = ("competencia", "matricula")
COMISSAO_KEY = ("competencia", "cod_marca", "cod_cargo")
MOVEMENT_FIELDS = ("cod_loja", "cod_marca", "cod_cargo", "descr_cargo")
REVIEW_AUTHORITY = "https://github.com/Titus-System/synapse/pull/86#pullrequestreview-5182423687"


@dataclass(frozen=True)
class CompetenciaSource:
    competencia: str
    rh_file: Path
    rh_sheet: str
    vendas_file: Path
    vendas_sheet: str


@dataclass(frozen=True)
class SourceRow:
    row_number: int
    values: dict[str, CellValue]


@dataclass(frozen=True)
class SourceTable:
    headers: tuple[str, ...]
    rows: tuple[SourceRow, ...]


@dataclass(frozen=True)
class CanonicalRow:
    source_file: Path
    source_sheet: str
    source_row: int
    record: JsonObject


@dataclass(frozen=True)
class BuildResult:
    output_dir: Path
    artifacts: tuple[Path, ...]
    hashes: dict[str, str]


def manifest(source_root: Path) -> tuple[CompetenciaSource, ...]:
    return (
        CompetenciaSource(
            competencia="2025-07",
            rh_file=source_root / "BASE RH" / "BASE RH_JUL25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_JUL25.xlsx",
            vendas_sheet="VENDAS",
        ),
        CompetenciaSource(
            competencia="2025-08",
            rh_file=source_root / "BASE RH" / "BASE RH_AGO25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_AGO25.xlsx",
            vendas_sheet="VENDAS",
        ),
        CompetenciaSource(
            competencia="2025-09",
            rh_file=source_root / "BASE RH" / "BASE RH_SET25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_SET25.xlsx",
            vendas_sheet="VENDAS",
        ),
        CompetenciaSource(
            competencia="2025-10",
            rh_file=source_root / "BASE RH" / "BASE RH_OUT25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_OUT25.xlsx",
            vendas_sheet="VENDAS",
        ),
        CompetenciaSource(
            competencia="2025-11",
            rh_file=source_root / "BASE RH" / "BASE RH_NOV25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_NOV25.xlsx",
            vendas_sheet="VENDAS",
        ),
        CompetenciaSource(
            competencia="2025-12",
            rh_file=source_root / "BASE RH" / "BASE RH_DEZ25.xlsx",
            rh_sheet="BASE HC",
            vendas_file=source_root / "BASE_VENDAS" / "BASE_VENDAS_DEZ25.xlsx",
            vendas_sheet="VENDAS",
        ),
    )


def build_dataset(
    source_root: Path = DEFAULT_SOURCE_ROOT,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
) -> BuildResult:
    report = empty_report()
    competencia_sources = manifest(source_root)
    source_rh: list[CanonicalRow] = []
    source_vendas: list[CanonicalRow] = []

    for source in competencia_sources:
        source_rh.extend(build_rh_rows(source, report))
        source_vendas.extend(build_vendas_rows(source, report))

    history = build_history(source_rh)
    register_history(history, report)
    rh_rows = fill_rh_continuity(history, reconcile_rh(history, source_rh, report), report)
    vendas_rows = [
        row for row in source_vendas if row.record["competencia"] in PUBLISHED_COMPETENCIAS
    ]
    vendas_rows = reconcile_sales(history, rh_rows, vendas_rows, report)
    comissoes_rows = build_comissoes_rows(
        source_root / "BASE_COMMISS_FINAL.xlsx",
        "Commission",
        PUBLISHED_COMPETENCIAS,
        report,
    )
    validate_commission_matches(rh_rows, comissoes_rows)
    register_dataset_diagnostics(source_rh, source_vendas, rh_rows, vendas_rows, history, report)
    report["checks"] = build_checks(rh_rows, vendas_rows, comissoes_rows, report)

    output_dir.mkdir(parents=True, exist_ok=True)
    artifacts = (
        output_dir / "rh.jsonl",
        output_dir / "vendas.jsonl",
        output_dir / "comissoes.jsonl",
        output_dir / "schema.json",
        output_dir / "normalization_report.json",
    )
    write_jsonl(artifacts[0], (row.record for row in rh_rows))
    write_jsonl(artifacts[1], (row.record for row in vendas_rows))
    write_jsonl(artifacts[2], (row.record for row in comissoes_rows))
    write_json(artifacts[3], build_schema())
    write_json(artifacts[4], report)

    return BuildResult(
        output_dir=output_dir,
        artifacts=artifacts,
        hashes={path.name: sha256_file(path) for path in artifacts},
    )


def empty_report() -> JsonObject:
    return {
        "version": 3,
        "source_validations": {
            "rh": {},
            "vendas": {},
            "comissoes": {},
        },
        "normalizations_applied": [],
        "discarded_rows": [],
        "discarded_columns": [],
        "reconciliations": [],
        "warnings": [],
        "invariants": [],
        "competency_status": {},
        "decisions": build_decisions(),
        "checks": {},
    }


def build_rh_rows(source: CompetenciaSource, report: JsonObject) -> list[CanonicalRow]:
    table = read_sheet(source.rh_file, source.rh_sheet)
    validate_columns(
        headers=table.headers,
        expected=tuple(RH_COLUMN_MAP),
        allowed_extra=("VENDAS R$",) if source.competencia == "2025-09" else (),
        source_file=source.rh_file,
    )
    validate_data_demiss_alignment(source, table)
    if "VENDAS R$" in table.headers:
        discarded_columns(report).append(
            {
                "dataset": "rh",
                "competencia": source.competencia,
                "source_file": relative_source(source.rh_file),
                "source_sheet": source.rh_sheet,
                "source_column": "VENDAS R$",
                "reason": "not_in_canonical_rh_schema",
            }
        )

    rows = [
        CanonicalRow(
            source_file=source.rh_file,
            source_sheet=source.rh_sheet,
            source_row=row.row_number,
            record={
                "competencia": source.competencia,
                "data_ref": required_date(row.values["Data_Ref"], "Data_Ref"),
                "cod_marca": required_int(row.values["Cod_Marca"], "Cod_Marca"),
                "descr_marca": required_string(row.values["Descri_Marca"], "Descri_Marca"),
                "cod_loja": required_int(row.values["Cod_Loja"], "Cod_Loja"),
                "descr_loja": required_string(row.values["Descr_Loja"], "Descr_Loja"),
                "matricula": required_string(row.values["Matricula"], "Matricula"),
                "data_admiss": required_date(row.values["Data_Admiss"], "Data_Admiss"),
                "data_demiss": optional_date(row.values["Data_Demiss"], "Data_Demiss"),
                "cod_cargo": required_int(row.values["Cod_Cargo"], "Cod_Cargo"),
                "descr_cargo": required_string(row.values["Descri_Cargo"], "Descri_Cargo"),
            },
        )
        for row in table.rows
    ]
    validate_reference_month(rows)
    rows = deduplicate_rh(rows, report)
    register_rh_validation(source, table, rows, report)
    register_date_normalizations(
        report,
        dataset="rh",
        competencia=source.competencia,
        fields=("Data_Ref", "Data_Admiss", "Data_Demiss"),
    )
    return rows


def build_vendas_rows(source: CompetenciaSource, report: JsonObject) -> list[CanonicalRow]:
    table = read_sheet(source.vendas_file, source.vendas_sheet)
    validate_columns(
        headers=table.headers,
        expected=tuple(VENDAS_COLUMN_MAP),
        allowed_extra=(),
        source_file=source.vendas_file,
    )
    rows = [build_venda_row(source, row) for row in table.rows]
    validate_reference_month(rows)
    register_vendas_validation(source, table, rows, report)
    register_date_normalizations(
        report,
        dataset="vendas",
        competencia=source.competencia,
        fields=("Date_Ref",),
    )
    return rows


def build_venda_row(source: CompetenciaSource, row: SourceRow) -> CanonicalRow:
    data_ref = required_date(row.values["Date_Ref"], "Date_Ref")
    return CanonicalRow(
        source_file=source.vendas_file,
        source_sheet=source.vendas_sheet,
        source_row=row.row_number,
        record={
            "competencia": source.competencia,
            "data_ref": data_ref,
            "data_venda": data_ref if is_real_sale_date(source.competencia, data_ref) else None,
            "cod_marca": required_int(row.values["Cod_Marca"], "Cod_Marca"),
            "descr_marca": required_string(row.values["Descr_Marca"], "Descr_Marca"),
            "cod_loja": required_int(row.values["Cod_Loja"], "Cod_Loja"),
            "descr_loja": required_string(row.values["Descr_Loja"], "Descr_Loja"),
            "matricula": required_string(row.values["Matricula"], "Matricula"),
            "vlr_venda": required_number(row.values["Vlr _Venda"], "Vlr _Venda"),
        },
    )


def build_comissoes_rows(
    source_file: Path,
    sheet_name: str,
    competencias: tuple[str, ...],
    report: JsonObject,
) -> list[CanonicalRow]:
    table = read_sheet(source_file, sheet_name)
    validate_columns(
        headers=table.headers,
        expected=tuple(COMISSOES_COLUMN_MAP),
        allowed_extra=(),
        source_file=source_file,
    )
    base_rows: list[CanonicalRow] = []
    source_key_counts: Counter[tuple[int, int]] = Counter()
    for row in table.rows:
        record: JsonObject = {
            "cod_marca": required_int(row.values["Cod_Marca"], "Cod_Marca"),
            "descr_marca": required_string(row.values["Descr_Marca"], "Descr_Marca"),
            "cod_cargo": required_int(row.values["Cod_Cargo"], "Cod_Cargo"),
            "descr_cargo": required_string(row.values["Descri_Cargo"], "Descri_Cargo"),
            "percentual_comissao": required_number(row.values["%_Comiss"], "%_Comiss"),
        }
        source_key_counts[(cast(int, record["cod_marca"]), cast(int, record["cod_cargo"]))] += 1
        if record["cod_cargo"] == 150 and record["descr_cargo"] == "GERENTE QUIOSQUE":
            discarded_rows(report).append(
                {
                    "dataset": "comissoes",
                    "source_file": relative_source(source_file),
                    "source_sheet": sheet_name,
                    "source_row": row.row_number,
                    **record,
                    "reason": "manager_kiosk_rate_superseded_by_canonical_manager_rate",
                }
            )
            continue
        if record["cod_cargo"] == 150 and record["descr_cargo"] != "GERENTE DE LOJA":
            raise ValueError(f"Unknown manager rate description: {record['descr_cargo']}")
        base_rows.append(CanonicalRow(source_file, sheet_name, row.row_number, record))
    rows = [
        CanonicalRow(
            row.source_file,
            row.source_sheet,
            row.source_row,
            {"competencia": competencia, **row.record},
        )
        for competencia in competencias
        for row in base_rows
    ]
    validate_commission_matches([], rows)
    source_validations(report)["comissoes"] = {
        "source_file": relative_source(source_file),
        "source_sheet": sheet_name,
        "source_row_count": len(table.rows),
        "canonical_row_count": len(rows),
        "canonical_rules_per_competencia": len(base_rows),
        "ambiguous_source_keys": [
            {"cod_marca": marca, "cod_cargo": cargo, "source_rows": count}
            for (marca, cargo), count in sorted(source_key_counts.items())
            if count > 1
        ],
        "expanded_competencias": list(competencias),
    }
    return rows


def read_sheet(path: Path, sheet_name: str) -> SourceTable:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        header_values = next(rows)
        headers = tuple(normalize_header(value) for value in header_values)
        if "" in headers:
            raise ValueError(f"{path} has an empty header")
        if len(set(headers)) != len(headers):
            raise ValueError(f"{path} has duplicate headers after trimming")

        source_rows: list[SourceRow] = []
        for row_number, row_values in enumerate(rows, start=2):
            values = tuple(cast(CellValue, value) for value in row_values[: len(headers)])
            if all(value is None for value in values):
                continue
            source_rows.append(SourceRow(row_number, dict(zip(headers, values, strict=True))))
        return SourceTable(headers, tuple(source_rows))
    finally:
        workbook.close()


def validate_columns(
    *,
    headers: tuple[str, ...],
    expected: tuple[str, ...],
    allowed_extra: tuple[str, ...],
    source_file: Path,
) -> None:
    missing = sorted(set(expected) - set(headers))
    extra = sorted(set(headers) - set(expected) - set(allowed_extra))
    if missing or extra:
        raise ValueError(f"{source_file} has invalid columns: missing={missing}, extra={extra}")


def validate_data_demiss_alignment(source: CompetenciaSource, table: SourceTable) -> None:
    for row in table.rows:
        data_demiss = row.values["Data_Demiss"]
        cargo_code = known_cargo_code_value(data_demiss)
        if cargo_code is not None:
            raise ValueError(
                "Possible Data_Demiss/Cod_Cargo misalignment: "
                f"file={source_label(source.rh_file)}, "
                f"competencia={source.competencia}, "
                f"row={row.row_number}, "
                f"Data_Demiss={cargo_code}"
            )


def known_cargo_code_value(value: CellValue) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int) and value in KNOWN_CARGO_CODES:
        return value
    if isinstance(value, float) and value.is_integer() and int(value) in KNOWN_CARGO_CODES:
        return int(value)
    return None


def deduplicate_rh(rows: list[CanonicalRow], report: JsonObject) -> list[CanonicalRow]:
    groups: dict[tuple[str, str], list[CanonicalRow]] = defaultdict(list)
    for row in rows:
        groups[(cast(str, row.record["competencia"]), cast(str, row.record["matricula"]))].append(
            row
        )
    result: list[CanonicalRow] = []
    for (competencia, matricula), duplicates in groups.items():
        kept = duplicates[0]
        if any(row.record != kept.record for row in duplicates[1:]):
            locations = [
                (source_label(row.source_file), row.source_sheet, row.source_row)
                for row in duplicates
            ]
            raise ValueError(
                f"Conflicting RH rows: competencia={competencia}, matricula={matricula}, "
                f"source_rows={locations}"
            )
        result.append(kept)
        for row in duplicates[1:]:
            discarded_rows(report).append(
                {
                    "dataset": "rh",
                    **provenance(row),
                    "matricula": matricula,
                    "reason": "duplicate_integral_rh_row",
                    "kept_source_row": kept.source_row,
                }
            )
    return result


def validate_reference_month(rows: list[CanonicalRow]) -> None:
    for row in rows:
        if cast(str, row.record["data_ref"])[:7] != row.record["competencia"]:
            raise ValueError(
                f"data_ref month differs from manifest: competencia={row.record['competencia']}, "
                f"data_ref={row.record['data_ref']}, file={source_label(row.source_file)}, "
                f"sheet={row.source_sheet}, row={row.source_row}"
            )


def register_rh_validation(
    source: CompetenciaSource,
    table: SourceTable,
    rows: list[CanonicalRow],
    report: JsonObject,
) -> None:
    source_validations(report)["rh"][source.competencia] = {
        "source_file": relative_source(source.rh_file),
        "source_sheet": source.rh_sheet,
        "source_row_count": len(table.rows),
        "normalized_row_count": len(rows),
        "canonical_row_count": len(rows) if source.competencia in PUBLISHED_COMPETENCIAS else 0,
        "headers_after_trim": list(table.headers),
        "data_demiss_cod_cargo_alignment": {
            "validation": "reject_invalid_values_before_normalization",
            "validated_source_rows": len(table.rows),
            "known_cargo_codes_checked": sorted(KNOWN_CARGO_CODES),
            "failure_behavior": "raise_before_artifact_writes",
        },
        "reference_month_validated_rows": len(rows),
    }


def register_vendas_validation(
    source: CompetenciaSource,
    table: SourceTable,
    rows: list[CanonicalRow],
    report: JsonObject,
) -> None:
    date_ref_counts: dict[str, int] = defaultdict(int)
    non_month_reference_dates: dict[str, int] = defaultdict(int)
    for row in rows:
        data_ref = cast(str, row.record["data_ref"])
        date_ref_counts[data_ref] += 1
        if not data_ref.endswith("-01"):
            non_month_reference_dates[data_ref] += 1

    source_validations(report)["vendas"][source.competencia] = {
        "source_file": relative_source(source.vendas_file),
        "source_sheet": source.vendas_sheet,
        "source_row_count": len(table.rows),
        "normalized_row_count": len(rows),
        "canonical_row_count": len(rows) if source.competencia in PUBLISHED_COMPETENCIAS else 0,
        "reference_month_validated_rows": len(rows),
        "date_ref_check": {
            "checked": True,
            "non_day_one_count": sum(non_month_reference_dates.values()),
            "non_day_one_dates": sorted_dict(non_month_reference_dates),
            "all_date_ref_counts": sorted_dict(date_ref_counts),
        },
        "source_total_vlr_venda": decimal_text(sum_source_vlr_venda(table)),
        "normalized_total_vlr_venda": decimal_text(sum_record_number(rows, "vlr_venda")),
        "canonical_total_vlr_venda": (
            decimal_text(sum_record_number(rows, "vlr_venda"))
            if source.competencia in PUBLISHED_COMPETENCIAS
            else "0"
        ),
    }
    if sum_source_vlr_venda(table) != sum_record_number(rows, "vlr_venda"):
        raise ValueError(f"Sales total differs from source: competencia={source.competencia}")


def build_checks(
    rh_rows: list[CanonicalRow],
    vendas_rows: list[CanonicalRow],
    comissoes_rows: list[CanonicalRow],
    report: JsonObject,
) -> JsonObject:
    sales_counts = Counter(cast(str, r.record["competencia"]) for r in vendas_rows)
    real_dates = Counter(
        cast(str, r.record["competencia"])
        for r in vendas_rows
        if r.record["data_venda"] is not None
    )
    commission_source = source_validations(report)["comissoes"]
    return {
        "determinism": {
            "encoding": "utf-8",
            "line_endings": "LF",
            "field_order": "fixed",
            "record_order": "manifest order; observed rows then synthetic RH by matricula",
            "timestamps": "not_included",
            "uuids": "not_included",
        },
        "published_competencias": list(PUBLISHED_COMPETENCIAS),
        "rh": {
            "total_rows": len(rh_rows),
            "rows_by_competencia": counts_for_competencias(
                dict(Counter(cast(str, r.record["competencia"]) for r in rh_rows)),
                PUBLISHED_COMPETENCIAS,
            ),
            "unique_key": list(RH_KEY),
            "admission_reconciliations": len(
                reconciliations_of(report, "canonical_admission_date")
            ),
            "termination_reconciliations": len(
                reconciliations_of(report, "terminal_termination_propagation")
            ),
        },
        "vendas": {
            "total_rows": len(vendas_rows),
            "rows_by_competencia": counts_for_competencias(
                dict(sales_counts), PUBLISHED_COMPETENCIAS
            ),
            "data_venda_rows_by_competencia": counts_for_competencias(
                dict(real_dates), PUBLISHED_COMPETENCIAS
            ),
            "duplicates_policy": "preserved",
        },
        "comissoes": {
            "source_rows": commission_source["source_row_count"],
            "rules_per_competencia": commission_source["canonical_rules_per_competencia"],
            "competencias": len(PUBLISHED_COMPETENCIAS),
            "expected_rows": cast(int, commission_source["canonical_rules_per_competencia"])
            * len(PUBLISHED_COMPETENCIAS),
            "actual_rows": len(comissoes_rows),
            "unique_key": list(COMISSAO_KEY),
            "rh_rows_with_exactly_one_match": len(rh_rows),
        },
    }


def register_date_normalizations(
    report: JsonObject,
    *,
    dataset: str,
    competencia: str,
    fields: tuple[str, ...],
) -> None:
    normalizations_applied(report).append(
        {
            "dataset": dataset,
            "competencia": competencia,
            "fields": list(fields),
            "normalization": "date_cells_to_iso_yyyy_mm_dd",
        }
    )


def is_real_sale_date(competencia: str, data_ref: str) -> bool:
    return competencia == "2025-11" and data_ref in {
        "2025-11-24",
        "2025-11-25",
        "2025-11-26",
        "2025-11-27",
        "2025-11-28",
    }


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def required_string(value: CellValue, field_name: str) -> str:
    if value is None:
        raise ValueError(f"{field_name} is required")
    return str(value)


def required_int(value: CellValue, field_name: str) -> int:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{field_name} must be an integer")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise ValueError(f"{field_name} must be an integer, got {value!r}")


def required_number(value: CellValue, field_name: str) -> int | float:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{field_name} must be numeric")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    raise ValueError(f"{field_name} must be numeric, got {value!r}")


def optional_date(value: CellValue, field_name: str) -> str | None:
    if value is None:
        return None
    return required_date(value, field_name)


def required_date(value: CellValue, field_name: str) -> str:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{field_name} must be a date")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float):
        return cast(datetime, from_excel(value)).date().isoformat()
    if isinstance(value, str):
        return date.fromisoformat(value).isoformat()
    raise ValueError(f"{field_name} must be a date, got {value!r}")


def sum_source_vlr_venda(table: SourceTable) -> Decimal:
    return sum(
        (Decimal(str(row.values["Vlr _Venda"])) for row in table.rows),
        Decimal("0"),
    )


def sum_record_number(rows: list[CanonicalRow], field_name: str) -> Decimal:
    return sum(
        (Decimal(str(row.record[field_name])) for row in rows),
        Decimal("0"),
    )


def decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def sorted_dict(values: dict[str, int]) -> JsonObject:
    return {key: values[key] for key in sorted(values)}


def counts_for_competencias(values: dict[str, int], competencias: Sequence[str]) -> JsonObject:
    return {competencia: values.get(competencia, 0) for competencia in competencias}


def relative_source(path: Path) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT.resolve()).as_posix()
    except ValueError:
        return f"external_sources/{path.parent.name}/{path.name}"


def source_label(path: Path) -> str:
    try:
        return relative_source(path)
    except ValueError:
        return str(path)


def write_jsonl(path: Path, rows: object) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as file:
        for row in cast(Any, rows):
            file.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")))
            file.write("\n")


def write_json(path: Path, data: JsonObject) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")


def sha256_file(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_validations(report: JsonObject) -> dict[str, JsonObject]:
    return cast(dict[str, JsonObject], report["source_validations"])


def normalizations_applied(report: JsonObject) -> list[JsonObject]:
    return cast(list[JsonObject], report["normalizations_applied"])


def discarded_rows(report: JsonObject) -> list[JsonObject]:
    return cast(list[JsonObject], report["discarded_rows"])


def discarded_columns(report: JsonObject) -> list[JsonObject]:
    return cast(list[JsonObject], report["discarded_columns"])


def warnings(report: JsonObject) -> list[JsonObject]:
    return cast(list[JsonObject], report["warnings"])


def provenance(row: CanonicalRow) -> JsonObject:
    return {
        "source_file": relative_source(row.source_file),
        "source_sheet": row.source_sheet,
        "source_row": row.source_row,
        "competencia": row.record["competencia"],
    }


def rh_evidence(row: CanonicalRow) -> JsonObject:
    return {**provenance(row), **row.record}


def build_history(rows: list[CanonicalRow]) -> dict[str, list[CanonicalRow]]:
    history: dict[str, list[CanonicalRow]] = defaultdict(list)
    for row in rows:
        history[cast(str, row.record["matricula"])].append(row)
    return {
        matricula: sorted(
            records,
            key=lambda r: (
                cast(str, r.record["competencia"]),
                r.source_file.as_posix(),
                r.source_row,
            ),
        )
        for matricula, records in sorted(history.items())
    }


def canonical_dates(records: list[CanonicalRow]) -> tuple[CanonicalRow, CanonicalRow | None]:
    admission = min(
        records,
        key=lambda r: (
            cast(str, r.record["data_admiss"]),
            cast(str, r.record["competencia"]),
            r.source_row,
        ),
    )
    terminations = [r for r in records if r.record["data_demiss"] is not None]
    termination = min(
        terminations,
        key=lambda r: (
            cast(str, r.record["data_demiss"]),
            cast(str, r.record["competencia"]),
            r.source_row,
        ),
        default=None,
    )
    return admission, termination


def register_history(history: dict[str, list[CanonicalRow]], report: JsonObject) -> None:
    admissions: list[JsonObject] = []
    for matricula, records in history.items():
        if len({r.record["data_admiss"] for r in records}) > 1:
            admission, _ = canonical_dates(records)
            admissions.append(
                {
                    "matricula": matricula,
                    "canonical_value": admission.record["data_admiss"],
                    "observations": [rh_evidence(r) for r in records],
                }
            )
        for previous, current in pairwise(records):
            before = cast(str, previous.record["competencia"])
            after = cast(str, current.record["competencia"])
            for field_name in MOVEMENT_FIELDS:
                if previous.record[field_name] != current.record[field_name]:
                    warnings(report).append(
                        {
                            "invariant_id": "I5",
                            "classification": "observed_attribute_change",
                            "matricula": matricula,
                            "field": field_name,
                            "previous_competencia": before,
                            "current_competencia": after,
                            "previous_value": previous.record[field_name],
                            "current_value": current.record[field_name],
                            "previous_source": provenance(previous),
                            "current_source": provenance(current),
                            "intermediate_missing_competencias": months_between(before, after),
                            "action": "preserved",
                            "effective_date": None,
                            "blocking": False,
                        }
                    )
    source_validations(report)["history"] = {
        "analyzed_competencias": sorted(
            {
                cast(str, row.record["competencia"])
                for records in history.values()
                for row in records
            }
        ),
        "admission_variations": admissions,
        "matriculas_with_admission_variations": len(admissions),
    }


def months_between(before: str, after: str) -> list[str]:
    year, month = map(int, before.split("-"))
    result: list[str] = []
    while True:
        month += 1
        if month == 13:
            year, month = year + 1, 1
        value = f"{year:04d}-{month:02d}"
        if value >= after:
            return result
        result.append(value)


def reconciliations_of(report: JsonObject, kind: str) -> list[JsonObject]:
    return [r for r in cast(list[JsonObject], report["reconciliations"]) if r["type"] == kind]


def reconcile_rh(
    history: dict[str, list[CanonicalRow]],
    source_rows: list[CanonicalRow],
    report: JsonObject,
) -> list[CanonicalRow]:
    dates = {matricula: canonical_dates(records) for matricula, records in history.items()}
    result: list[CanonicalRow] = []
    reconciliations = cast(list[JsonObject], report["reconciliations"])
    for row in source_rows:
        competencia = cast(str, row.record["competencia"])
        if competencia not in PUBLISHED_COMPETENCIAS:
            continue
        record = row.record.copy()
        matricula = cast(str, record["matricula"])
        admission, termination = dates[matricula]
        if record["data_admiss"] != admission.record["data_admiss"]:
            reconciliations.append(
                {
                    "type": "canonical_admission_date",
                    "matricula": matricula,
                    **provenance(row),
                    "original_value": record["data_admiss"],
                    "canonical_value": admission.record["data_admiss"],
                    "evidence_competencia": admission.record["competencia"],
                    "evidence_source_file": relative_source(admission.source_file),
                    "evidence_source_sheet": admission.source_sheet,
                    "evidence_source_row": admission.source_row,
                }
            )
            record["data_admiss"] = admission.record["data_admiss"]
        if termination is not None:
            terminal_date = cast(str, termination.record["data_demiss"])
            if competencia >= terminal_date[:7] and record["data_demiss"] != terminal_date:
                reconciliations.append(
                    {
                        "type": "terminal_termination_propagation",
                        "matricula": matricula,
                        **provenance(row),
                        "origin_competencia": termination.record["competencia"],
                        "origin_data_demiss": terminal_date,
                        "target_competencia": competencia,
                        "original_value": record["data_demiss"],
                        "canonical_value": terminal_date,
                        "evidence_source_file": relative_source(termination.source_file),
                        "evidence_source_sheet": termination.source_sheet,
                        "evidence_source_row": termination.source_row,
                    }
                )
                record["data_demiss"] = terminal_date
        result.append(CanonicalRow(row.source_file, row.source_sheet, row.source_row, record))
    return result


def eligible_rh_month(records: list[CanonicalRow], competencia: str) -> bool:
    admission, termination = canonical_dates(records)
    return cast(str, admission.record["data_admiss"])[:7] <= competencia and (
        termination is None or cast(str, termination.record["data_demiss"])[:7] >= competencia
    )


def fill_rh_continuity(
    history: dict[str, list[CanonicalRow]], observed: list[CanonicalRow], report: JsonObject
) -> list[CanonicalRow]:
    keys = {(r.record["competencia"], r.record["matricula"]) for r in observed}
    synthetic: list[CanonicalRow] = []
    for competencia in PUBLISHED_COMPETENCIAS:
        for matricula, records in sorted(history.items()):
            if (competencia, matricula) in keys or not eligible_rh_month(records, competencia):
                continue
            previous = [r for r in records if cast(str, r.record["competencia"]) < competencia]
            following = [r for r in records if cast(str, r.record["competencia"]) > competencia]
            base = previous[-1] if previous else following[0]
            admission, termination = canonical_dates(records)
            terminal_date = cast(str, termination.record["data_demiss"]) if termination else None
            record = base.record.copy()
            record.update(
                {
                    "competencia": competencia,
                    "data_ref": competencia + "-01",
                    "data_admiss": admission.record["data_admiss"],
                    "data_demiss": terminal_date
                    if terminal_date and competencia >= terminal_date[:7]
                    else None,
                }
            )
            synthetic.append(
                CanonicalRow(base.source_file, base.source_sheet, base.source_row, record)
            )
            cast(list[JsonObject], report["reconciliations"]).append(
                {
                    "type": "rh_continuity_fill",
                    "matricula": matricula,
                    "competencia": competencia,
                    "source_direction": "previous" if previous else "next",
                    "evidence_competencia": base.record["competencia"],
                    "evidence_source_file": relative_source(base.source_file),
                    "evidence_source_sheet": base.source_sheet,
                    "evidence_source_row": base.source_row,
                    "canonical_admission_date": admission.record["data_admiss"],
                    "terminal_termination_date": terminal_date,
                    "reason": "maintain_employee_continuity",
                }
            )
    return [
        row
        for competencia in PUBLISHED_COMPETENCIAS
        for row in observed + synthetic
        if row.record["competencia"] == competencia
    ]


def sale_after_termination(row: CanonicalRow, terminal_date: str | None) -> bool:
    if terminal_date is None:
        return False
    actual_date = cast(str | None, row.record["data_venda"])
    if actual_date is not None:
        return actual_date > terminal_date
    return cast(str, row.record["competencia"]) > terminal_date[:7]


def reconcile_sales(
    history: dict[str, list[CanonicalRow]],
    rh: list[CanonicalRow],
    sales: list[CanonicalRow],
    report: JsonObject,
) -> list[CanonicalRow]:
    keys = {(r.record["competencia"], r.record["matricula"]) for r in rh}
    terminals = {
        matricula: cast(str, termination.record["data_demiss"]) if termination else None
        for matricula, records in history.items()
        for _, termination in [canonical_dates(records)]
    }
    result: list[CanonicalRow] = []
    for row in sales:
        matricula = cast(str, row.record["matricula"])
        terminal_date = terminals.get(matricula)
        reason = None
        if sale_after_termination(row, terminal_date):
            reason = "sale_after_terminal_termination"
        elif (row.record["competencia"], matricula) not in keys:
            reason = "sale_without_reconstructable_rh"
        if reason:
            discarded_rows(report).append(
                {
                    "dataset": "vendas",
                    **provenance(row),
                    **{
                        field: row.record[field]
                        for field in (
                            "matricula",
                            "data_ref",
                            "data_venda",
                            "cod_marca",
                            "cod_loja",
                            "vlr_venda",
                        )
                    },
                    "terminal_termination_date": terminal_date,
                    "reason": reason,
                }
            )
        else:
            result.append(row)
    return result


def commission_key(row: CanonicalRow) -> tuple[str, int, int]:
    return (
        cast(str, row.record["competencia"]),
        cast(int, row.record["cod_marca"]),
        cast(int, row.record["cod_cargo"]),
    )


def validate_commission_matches(rh: list[CanonicalRow], commissions: list[CanonicalRow]) -> None:
    keys = Counter(commission_key(row) for row in commissions)
    duplicates = [key for key, count in sorted(keys.items()) if count != 1]
    if duplicates:
        raise ValueError(f"Commission keys must be unique: {duplicates}")
    for row in rh:
        key = commission_key(row)
        if keys[key] != 1:
            raise ValueError(
                f"Expected exactly one commission: key={key}, matricula={row.record['matricula']}"
            )


def build_decisions() -> list[JsonObject]:
    definitions = [
        (
            "EXCLUDE_2025_07",
            "unreconcilable_source_population",
            "July has a partial employee population and incompatible sales scale. "
            "Read it as historical evidence; exclude it from published tables.",
            ["2025-07"],
        ),
        (
            "CANONICAL_ADMISSION_MINIMUM",
            "minimum_observed_admission",
            "Use the minimum admission date across all six RH sources in every published RH row.",
            list(PUBLISHED_COMPETENCIAS),
        ),
        (
            "TERMINATION_IS_TERMINAL",
            "first_known_termination",
            "Keep subsequent RH rows and propagate the earliest known termination. "
            "A termination before the competency makes the employee ineligible for commission.",
            list(PUBLISHED_COMPETENCIAS),
        ),
        (
            "CANONICAL_MANAGER_RATE",
            "manager_store_rate_is_canonical",
            "Use GERENTE DE LOJA rates for code 150; keep both descriptions in RH.",
            list(PUBLISHED_COMPETENCIAS),
        ),
    ]
    decisions: list[JsonObject] = [
        {
            "id": ident,
            "status": "defined",
            "reason": reason,
            "description": description,
            "authority": REVIEW_AUTHORITY,
            "affected_scope": scope,
        }
        for ident, reason, description, scope in definitions
    ]
    for ident, description in [
        ("RECONSTRUCT_MISSING_RH", "Reconstruct eligible missing RH using observed history."),
        ("SALE_STORE_IS_AUTHORITATIVE", "Preserve sale store and brand, including shared sales."),
        ("DISCARD_POST_TERMINATION_SALES", "Discard sales after terminal termination."),
        (
            "FILL_RH_CONTINUITY",
            "Fill eligible months from previous observed RH, else next observed RH.",
        ),
        ("DISCARD_UNRESOLVED_ORPHAN_SALES", "Discard sales without reconstructable RH."),
    ]:
        decisions.append(
            {
                "id": ident,
                "status": "defined",
                "description": description,
                "authority": "T-026 final policy instructions",
                "affected_scope": list(PUBLISHED_COMPETENCIAS),
            }
        )
    return decisions


def invariant(
    ident: str,
    name: str,
    unit: str,
    found: int,
    reconciled: int,
    remaining: int,
    details: list[JsonObject],
    *,
    status: str,
    blocking: bool = False,
) -> JsonObject:
    return {
        "id": ident,
        "name": name,
        "status": status,
        "blocking": blocking,
        "count_unit": unit,
        "violations_found": found,
        "reconciliations_applied": reconciled,
        "remaining_violations": remaining,
        "details": details,
    }


def dimension_diagnostics(rh: list[CanonicalRow], sales: list[CanonicalRow]) -> list[JsonObject]:
    index = {(r.record["competencia"], r.record["matricula"]): r for r in rh}
    by_month: dict[str, list[CanonicalRow]] = defaultdict(list)
    for row in sales:
        by_month[cast(str, row.record["competencia"])].append(row)
    result: list[JsonObject] = []
    for competencia, rows in sorted(by_month.items()):
        mismatches: list[JsonObject] = []
        counts: Counter[str] = Counter({"only_store": 0, "only_brand": 0, "store_and_brand": 0})
        dimensions: dict[str, set[tuple[int, int]]] = defaultdict(set)
        for row in rows:
            matricula = cast(str, row.record["matricula"])
            dimensions[matricula].add(
                (cast(int, row.record["cod_marca"]), cast(int, row.record["cod_loja"]))
            )
            employee = index.get((competencia, matricula))
            if employee is None:
                continue
            store_diff = row.record["cod_loja"] != employee.record["cod_loja"]
            brand_diff = row.record["cod_marca"] != employee.record["cod_marca"]
            if not (store_diff or brand_diff):
                continue
            kind = (
                "store_and_brand"
                if store_diff and brand_diff
                else ("only_store" if store_diff else "only_brand")
            )
            counts[kind] += 1
            mismatches.append(
                {
                    **provenance(row),
                    "matricula": matricula,
                    "type": kind,
                    "sale_cod_loja": row.record["cod_loja"],
                    "sale_cod_marca": row.record["cod_marca"],
                    "rh_cod_loja": employee.record["cod_loja"],
                    "rh_cod_marca": employee.record["cod_marca"],
                    "rh_source": provenance(employee),
                }
            )
        result.append(
            {
                "competencia": competencia,
                "sale_rows": len(mismatches),
                "by_type": dict(counts),
                "matriculas": sorted({cast(str, r["matricula"]) for r in mismatches}),
                "mismatches": mismatches,
                "multiple_store_sales": [
                    {
                        "matricula": matricula,
                        "dimensions": [
                            {"cod_marca": marca, "cod_loja": loja} for marca, loja in sorted(values)
                        ],
                    }
                    for matricula, values in sorted(dimensions.items())
                    if len({loja for _, loja in values}) > 1
                ],
            }
        )
    return result


def register_excluded_competency(
    source_rh: list[CanonicalRow],
    source_sales: list[CanonicalRow],
    report: JsonObject,
) -> JsonObject:
    july = {
        cast(str, r.record["matricula"]): r
        for r in source_rh
        if r.record["competencia"] == "2025-07"
    }
    august = {
        cast(str, r.record["matricula"]): r
        for r in source_rh
        if r.record["competencia"] == "2025-08"
    }
    missing = sorted(august.keys() - july.keys())
    old_missing = [m for m in missing if cast(str, august[m].record["data_admiss"]) < "2025-07-01"]
    comparison: list[JsonObject] = []
    for competencia in sorted({cast(str, r.record["competencia"]) for r in source_rh}):
        employees = [r for r in source_rh if r.record["competencia"] == competencia]
        sales = [r for r in source_sales if r.record["competencia"] == competencia]
        comparison.append(
            {
                "competencia": competencia,
                "rh_rows": len(employees),
                "sale_rows": len(sales),
                "rh_stores": len({r.record["cod_loja"] for r in employees}),
                "sale_stores": len({r.record["cod_loja"] for r in sales}),
                "roles": {
                    str(k): v
                    for k, v in sorted(
                        Counter(cast(int, r.record["cod_cargo"]) for r in employees).items()
                    )
                },
                "total_vlr_venda": decimal_text(sum_record_number(sales, "vlr_venda")),
                "maximum_sale": max(
                    (cast(int | float, r.record["vlr_venda"]) for r in sales), default=0
                ),
            }
        )
    detail: JsonObject = {
        "competencia": "2025-07",
        "status": "excluded",
        "published": False,
        "reason": "unreconcilable_source_population",
        "decision_id": "EXCLUDE_2025_07",
        "historical_evidence": True,
        "description": "Partial population and incompatible sales scale; no July reconstruction.",
        "august_matriculas_missing_in_july": missing,
        "admitted_before_july_missing_in_july": old_missing,
        "comparison": comparison,
    }
    cast(dict[str, JsonObject], report["competency_status"])["2025-07"] = detail
    for dataset, rows in [("rh", source_rh), ("vendas", source_sales)]:
        excluded = [r for r in rows if r.record["competencia"] == "2025-07"]
        detail[f"excluded_{dataset}_rows"] = len(excluded)
        detail[f"excluded_{dataset}_sources"] = (
            [
                {
                    "source_file": relative_source(excluded[0].source_file),
                    "source_sheet": excluded[0].source_sheet,
                    "first_source_row": min(r.source_row for r in excluded),
                    "last_source_row": max(r.source_row for r in excluded),
                    "row_count": len(excluded),
                }
            ]
            if excluded
            else []
        )
    return detail


def sales_reconciliation(
    original: list[CanonicalRow], final: list[CanonicalRow], discarded: list[JsonObject]
) -> JsonObject:
    source_total = sum_record_number(original, "vlr_venda")
    canonical_total = sum_record_number(final, "vlr_venda")
    discarded_total = sum((Decimal(str(r["vlr_venda"])) for r in discarded), Decimal(0))
    difference = source_total - discarded_total - canonical_total
    if len(original) - len(discarded) != len(final) or difference:
        raise ValueError("Published source minus discarded sales must equal canonical sales")
    return {
        "source_row_count": len(original),
        "discarded_row_count": len(discarded),
        "canonical_row_count": len(final),
        "source_total_vlr_venda": decimal_text(source_total),
        "discarded_total_vlr_venda": decimal_text(discarded_total),
        "canonical_total_vlr_venda": decimal_text(canonical_total),
        "reconciliation_difference": decimal_text(difference),
    }


def register_dataset_diagnostics(
    source_rh: list[CanonicalRow],
    source_sales: list[CanonicalRow],
    rh: list[CanonicalRow],
    sales: list[CanonicalRow],
    history: dict[str, list[CanonicalRow]],
    report: JsonObject,
) -> None:
    original = [r for r in source_sales if r.record["competencia"] in PUBLISHED_COMPETENCIAS]
    observed = [r for r in source_rh if r.record["competencia"] in PUBLISHED_COMPETENCIAS]
    keys = {(r.record["competencia"], r.record["matricula"]) for r in rh}
    if len(keys) != len(rh):
        raise ValueError("RH uniqueness violated after reconciliation")
    if any((r.record["competencia"], r.record["matricula"]) not in keys for r in sales):
        raise ValueError("Final sales must have exactly one RH match")
    missing_slots = [
        {"competencia": c, "matricula": m}
        for c in PUBLISHED_COMPETENCIAS
        for m, records in sorted(history.items())
        if eligible_rh_month(records, c) and (c, m) not in keys
    ]
    if missing_slots:
        raise ValueError("Eligible RH continuity is incomplete")
    for row in rh:
        admission, termination = canonical_dates(history[cast(str, row.record["matricula"])])
        if row.record["data_admiss"] != admission.record["data_admiss"]:
            raise ValueError("Canonical admission violated")
        if termination:
            value = cast(str, termination.record["data_demiss"])
            if (
                cast(str, row.record["competencia"]) >= value[:7]
                and row.record["data_demiss"] != value
            ):
                raise ValueError("Terminal termination violated")
    for row in sales:
        records = history[cast(str, row.record["matricula"])]
        _, termination = canonical_dates(records)
        terminal_date = cast(str, termination.record["data_demiss"]) if termination else None
        if sale_after_termination(row, terminal_date):
            raise ValueError("Final sales include post-termination sale")
    originals = {(r.source_file, r.source_sheet, r.source_row): r.record for r in original}
    if any(r.record != originals[(r.source_file, r.source_sheet, r.source_row)] for r in sales):
        raise ValueError("Retained sale values or dimensions changed")
    discarded = [r for r in discarded_rows(report) if r["dataset"] == "vendas"]
    discarded_index = {(r["source_file"], r["source_sheet"], r["source_row"]): r for r in discarded}
    if len(discarded_index) != len(discarded):
        raise ValueError("Sale discarded more than once")
    observed_keys = {(r.record["competencia"], r.record["matricula"]) for r in observed}
    orphans: list[JsonObject] = []
    for row in original:
        if (row.record["competencia"], row.record["matricula"]) in observed_keys:
            continue
        discard = discarded_index.get(
            (relative_source(row.source_file), row.source_sheet, row.source_row)
        )
        orphans.append(
            {
                **provenance(row),
                "matricula": row.record["matricula"],
                "action": discard["reason"] if discard else "reconstructed_rh",
                "blocking": False,
            }
        )
    dimensions = dimension_diagnostics(source_rh, source_sales)
    final_dimensions = dimension_diagnostics(rh, sales)
    source_validations(report)["sale_dimensions"] = {
        "by_competencia": dimensions,
        "final_by_competencia": final_dimensions,
        "blocking": False,
        "policy": "sale_store_and_brand_are_authoritative",
    }
    admissions = reconciliations_of(report, "canonical_admission_date")
    terminations = reconciliations_of(report, "terminal_termination_propagation")
    fills = reconciliations_of(report, "rh_continuity_fill")
    post_termination = [r for r in discarded if r["reason"] == "sale_after_terminal_termination"]
    movements = [
        w for w in warnings(report) if w.get("classification") == "observed_attribute_change"
    ]
    register_excluded_competency(source_rh, source_sales, report)
    totals: list[JsonObject] = []
    rh_counts: list[JsonObject] = []
    for competencia in PUBLISHED_COMPETENCIAS:
        detail = {
            "competencia": competencia,
            **sales_reconciliation(
                [r for r in original if r.record["competencia"] == competencia],
                [r for r in sales if r.record["competencia"] == competencia],
                [r for r in discarded if r["competencia"] == competencia],
            ),
        }
        totals.append(detail)
        cast(JsonObject, source_validations(report)["vendas"][competencia]).update(detail)
        rh_detail: JsonObject = {
            "competencia": competencia,
            "observed_row_count": sum(r.record["competencia"] == competencia for r in observed),
            "synthetic_row_count": sum(r["competencia"] == competencia for r in fills),
            "canonical_row_count": sum(r.record["competencia"] == competencia for r in rh),
        }
        rh_counts.append(rh_detail)
        cast(JsonObject, source_validations(report)["rh"][competencia]).update(rh_detail)
    report["financial_reconciliation"] = {
        "by_competencia": totals,
        "aggregate": sales_reconciliation(original, sales, discarded),
    }
    report["sale_discard_summary"] = {
        "by_reason": [
            {
                "reason": reason,
                "row_count": sum(r["reason"] == reason for r in discarded),
                "total_vlr_venda": decimal_text(
                    sum(
                        (Decimal(str(r["vlr_venda"])) for r in discarded if r["reason"] == reason),
                        Decimal(0),
                    )
                ),
            }
            for reason in ("sale_after_terminal_termination", "sale_without_reconstructable_rh")
        ],
        "dated_same_termination_month_discarded": sum(
            r["data_venda"] is not None
            and r["competencia"] == cast(str, r["terminal_termination_date"])[:7]
            for r in post_termination
        ),
        "undated_same_termination_month_preserved": sum(
            row.record["data_venda"] is None
            and termination is not None
            and row.record["competencia"] == cast(str, termination.record["data_demiss"])[:7]
            for row in sales
            for _, termination in [canonical_dates(history[cast(str, row.record["matricula"])])]
        ),
        "former_orphans_recovered": sum(r["action"] == "reconstructed_rh" for r in orphans),
        "former_orphans_discarded": sum(r["action"] != "reconstructed_rh" for r in orphans),
    }
    report["rh_population"] = {
        "by_competencia": rh_counts,
        "observed_row_count": len(observed),
        "synthetic_row_count": len(fills),
        "canonical_row_count": len(rh),
        "fills_by_direction": dict(Counter(cast(str, r["source_direction"]) for r in fills)),
    }
    validate_reference_month(rh + sales)
    duplicates = [
        r
        for r in discarded_rows(report)
        if r["reason"] == "duplicate_integral_rh_row" and r["competencia"] in PUBLISHED_COMPETENCIAS
    ]
    commission_details = [
        {"competencia": c, **key}
        for c in PUBLISHED_COMPETENCIAS
        for key in cast(
            list[JsonObject], source_validations(report)["comissoes"]["ambiguous_source_keys"]
        )
    ]
    report["invariants"] = [
        invariant(
            "I1",
            "unique_rh_per_competency",
            "rh_row",
            len(duplicates),
            len(duplicates),
            0,
            duplicates,
            status="reconciled" if duplicates else "passed",
        ),
        invariant(
            "I2",
            "sale_employee_exists_in_rh",
            "sale_row",
            len(orphans),
            len(orphans),
            0,
            orphans,
            status="reconciled" if orphans else "passed",
        ),
        invariant(
            "I3",
            "source_sale_dimensions_preserved",
            "sale_row",
            0,
            0,
            0,
            final_dimensions,
            status="passed",
        ),
        invariant(
            "I4",
            "termination_is_terminal",
            "rh_row_or_sale_row",
            len(terminations) + len(post_termination),
            len(terminations) + len(post_termination),
            0,
            [
                {
                    "rh_reconciliations": terminations,
                    "rh_reconciliation_count": len(terminations),
                    "discarded_sale_count": len(post_termination),
                    "discarded_sales": post_termination,
                }
            ],
            status="reconciled" if terminations or post_termination else "passed",
        ),
        invariant(
            "I5",
            "canonical_admission_and_recorded_attribute_changes",
            "rh_admission_field",
            len(admissions),
            len(admissions),
            0,
            [{"admission_reconciliations": admissions, "observed_attribute_changes": movements}],
            status="warning" if movements else ("reconciled" if admissions else "passed"),
        ),
        invariant(
            "I6",
            "exactly_one_commission_rate",
            "published_commission_key",
            len(commission_details),
            len(commission_details),
            0,
            commission_details,
            status="reconciled" if commission_details else "passed",
        ),
        invariant(
            "I7",
            "eligible_employee_month_continuity",
            "eligible_rh_month",
            len(fills),
            len(fills),
            0,
            fills,
            status="reconciled" if fills else "passed",
        ),
        invariant(
            "I8",
            "source_minus_discarded_equals_canonical",
            "competencia",
            0,
            0,
            0,
            totals,
            status="passed",
        ),
        invariant(
            "I9",
            "reference_month_matches_manifest",
            "rh_or_sale_row",
            0,
            0,
            0,
            [
                {
                    "competencia": c,
                    "validated_rows": sum(r.record["competencia"] == c for r in rh + sales),
                }
                for c in PUBLISHED_COMPETENCIAS
            ],
            status="passed",
        ),
    ]
    statuses = cast(dict[str, JsonObject], report["competency_status"])
    for competencia in PUBLISHED_COMPETENCIAS:
        statuses[competencia] = {
            "published": True,
            "status": "ready",
            "blocking_invariants": [],
            "blocking_decisions": [],
            "publication_semantics": "all_final_preparation_invariants_validated",
        }


def build_schema() -> JsonObject:
    return {
        "version": 3,
        "format": "jsonl",
        "published_competencias": list(PUBLISHED_COMPETENCIAS),
        "historical_evidence": {
            "excluded_competencias": ["2025-07"],
            "semantics": "July is read for history and date reconciliation, never published.",
        },
        "readiness": "Consult normalization_report.competency_status before simulation; "
        "ready means preparation invariants pass; movements are nonblocking observations.",
        "tables": {
            "baseline": {
                "files": [f"baselines/baseline-{c}.jsonl" for c in PUBLISHED_COMPETENCIAS],
                "produced_by": "python -m scripts.build_baselines (T-032, após T-026)",
                "schema_version": 1,
                "semantics": "Baseline não é gabarito; apuração interna da regra vigente.",
                "levels": {
                    "total": "Uma linha por competência, com total e asserções.",
                    "loja": "Uma linha por competência/cod_loja; comissão pela lotação no RH.",
                    "matricula": "Uma linha por competência/matricula; detalhes auditáveis.",
                },
                "aggregation": "Filtrar nivel antes de somar; os três níveis são a mesma apuração.",
                "traceability": "Linhas 1-based de rh/vendas completos; IDs de eventos e regras.",
                "fields": [
                    field("competencia", None, "string", False, "YYYY-MM", "Competência apurada."),
                    field("nivel", None, "string", False, None, "total, loja ou matricula."),
                    field(
                        "matricula", None, "string", True, None, "Preenchida no nível matricula."
                    ),
                    field("cod_loja", None, "integer", True, None, "Nulo no nível total."),
                    field("loja", None, "string", True, None, "Descrição da lotação no RH."),
                    field("cargo", None, "integer", True, None, "Preenchido no nível matricula."),
                    field("base_calculo", None, "number", True, None, "Base individual ajustada."),
                    field("comissao", None, "number", False, None, "BRL, duas casas decimais."),
                    field(
                        "rastreabilidade",
                        None,
                        "object",
                        True,
                        None,
                        "Fontes do cálculo individual.",
                    ),
                    field(
                        "assercoes",
                        None,
                        "array",
                        True,
                        None,
                        "Três desfechos T-031 na linha total.",
                    ),
                ],
            },
            "eventos_rh": {
                "file": "eventos_rh.jsonl",
                "produced_by": "python -m scripts.build_baselines",
                "primary_key": ["id"],
                "semantics": "T-028 preservada: mesmas datas/IDs e intervalos fechados.",
                "fields": [
                    field("id", None, "string", False, None, "ID da fonte T-028."),
                    field("tipo", None, "string", False, None, "Tipo do evento."),
                    field("matricula", None, "string", False, None, "Pessoa afetada."),
                    field(
                        "competencia_origem", None, "string", False, "YYYY-MM", "Mês do registro."
                    ),
                    field("data_inicio", None, "string", True, "YYYY-MM-DD", "Início inclusivo."),
                    field("data_fim", None, "string", True, "YYYY-MM-DD", "Fim inclusivo."),
                    field(
                        "detalhes",
                        None,
                        "object",
                        True,
                        None,
                        "Metadados originais, inclusive estimativas.",
                    ),
                ],
            },
            "regras_competencia": {
                "file": "regras_competencia.jsonl",
                "primary_key": ["id"],
                "semantics": "Catálogo histórico versionado de T-032; percentuais são frações.",
                "documentation": "worker/docs/t032-baselines.md",
                "fields": [
                    field("id", None, "string", False, None, "Referência única da política."),
                    field("competencia", None, "string", False, "YYYY-MM", "Mês de vigência."),
                    field("tipo", None, "string", False, None, "Operação determinística."),
                    field("fonte", None, "string", False, None, "Item da especificação."),
                ],
                "optional_fields_by_type": {
                    "selectors": ["marcas", "cargos", "cargos_excluidos", "matriculas"],
                    "bonus_final": ["valor", "admissao_ate"],
                    "bonus_base": ["valor"],
                    "substituir_percentual": ["valor"],
                    "copiar_percentual": ["marca_origem"],
                    "adicional_percentual": ["valor"],
                    "adicional_periodo": ["valor", "inicio", "fim"],
                    "bonus_faixa_individual": ["faixas"],
                    "bonus_faixa_loja": ["faixas"],
                },
            },
            "rh": {
                "file": "rh.jsonl",
                "continuity": {
                    "eligibility": "admission <= month end; termination absent or >= month start",
                    "evidence": "latest previous observed RH, otherwise nearest next observed RH",
                    "synthetic_evidence": False,
                    "post_termination": "retain observed; never synthesize after termination month",
                },
                "primary_key": list(RH_KEY),
                "unique_by": list(RH_KEY),
                "reference_month_constraint": "data_ref[:7] == competencia",
                "commission_join": {
                    "table": "comissoes",
                    "fields": list(COMISSAO_KEY),
                    "cardinality": "exactly_one",
                    "descr_cargo_participates": False,
                },
                "fields": [
                    field(
                        "competencia",
                        None,
                        "string",
                        False,
                        "YYYY-MM",
                        "Explicit source manifest competency.",
                    ),
                    field(
                        "data_ref",
                        "Data_Ref",
                        "string",
                        False,
                        "YYYY-MM-DD",
                        "RH source reference date; target month day one for reconstructed rows.",
                    ),
                    field("cod_marca", "Cod_Marca", "integer", False, None, "Brand code."),
                    field(
                        "descr_marca",
                        "Descri_Marca",
                        "string",
                        False,
                        None,
                        "Brand description.",
                    ),
                    field(
                        "cod_loja",
                        "Cod_Loja",
                        "integer",
                        False,
                        None,
                        "RH assignment in this competency; may differ from sale location.",
                    ),
                    field("descr_loja", "Descr_Loja", "string", False, None, "Store description."),
                    field(
                        "matricula",
                        "Matricula",
                        "string",
                        False,
                        None,
                        "Employee identifier preserved as text.",
                    ),
                    field(
                        "data_admiss",
                        "Data_Admiss",
                        "string",
                        False,
                        "YYYY-MM-DD",
                        "Minimum admission date across all six RH sources, including July; "
                        "applied to every published record, including earlier competencies.",
                    ),
                    field(
                        "data_demiss",
                        "Data_Demiss",
                        "string",
                        True,
                        "YYYY-MM-DD",
                        "Earliest known termination is terminal and propagated to subsequent RH. "
                        "A date before the competency starts makes the employee ineligible for "
                        "commission in that competency. Observed RH remains; "
                        "post-termination sales are discarded.",
                    ),
                    field("cod_cargo", "Cod_Cargo", "integer", False, None, "Role code."),
                    field(
                        "descr_cargo",
                        "Descri_Cargo",
                        "string",
                        False,
                        None,
                        "Role description.",
                    ),
                ],
            },
            "vendas": {
                "file": "vendas.jsonl",
                "discard_policy": {
                    "priority": [
                        "sale_after_terminal_termination",
                        "sale_without_reconstructable_rh",
                    ],
                    "dated_sale": "discard if data_venda > terminal termination date",
                    "undated_sale": "discard only if competencia > terminal termination month",
                    "financial_invariant": "source - discarded = canonical, exact Decimal",
                },
                "dimension_semantics": "Store and brand are the observed sale dimensions, "
                "preserved independently of RH assignment; sale store and brand are authoritative.",
                "reference_month_constraint": "data_ref[:7] == competencia",
                "fields": [
                    field(
                        "competencia",
                        None,
                        "string",
                        False,
                        "YYYY-MM",
                        "Compet\u00eancia expl\u00edcita do registro. Deve ser usada para "
                        "selecionar/agrupar compet\u00eancia e nunca deve ser inferida "
                        "de data_ref.",
                    ),
                    field(
                        "data_ref",
                        "Date_Ref",
                        "string",
                        False,
                        "YYYY-MM-DD",
                        "Normalized Date_Ref value from the source.",
                    ),
                    field(
                        "data_venda",
                        None,
                        "string",
                        True,
                        "YYYY-MM-DD",
                        "Data real da venda quando disponibilizada pela fonte; null quando "
                        "Date_Ref representa somente a refer\u00eancia mensal.",
                    ),
                    field("cod_marca", "Cod_Marca", "integer", False, None, "Brand code."),
                    field(
                        "descr_marca",
                        "Descr_Marca",
                        "string",
                        False,
                        None,
                        "Brand description.",
                    ),
                    field("cod_loja", "Cod_Loja", "integer", False, None, "Store code."),
                    field("descr_loja", "Descr_Loja", "string", False, None, "Store description."),
                    field(
                        "matricula",
                        "Matricula",
                        "string",
                        False,
                        None,
                        "Employee identifier preserved as text.",
                    ),
                    field(
                        "vlr_venda",
                        "Vlr _Venda",
                        "number",
                        False,
                        None,
                        "Source sale value preserved without two-decimal rounding.",
                    ),
                ],
            },
            "comissoes": {
                "file": "comissoes.jsonl",
                "primary_key": list(COMISSAO_KEY),
                "unique_by": list(COMISSAO_KEY),
                "manager_rate": {
                    "cod_cargo": 150,
                    "rh_descriptions": ["GERENTE DE LOJA", "GERENTE QUIOSQUE"],
                    "source_description": "GERENTE DE LOJA",
                    "semantics": "Use the store manager rate of the corresponding brand for "
                    "both descriptions. descr_cargo is descriptive, not a join field.",
                    "decision_id": "CANONICAL_MANAGER_RATE",
                },
                "fields": [
                    field(
                        "competencia",
                        None,
                        "string",
                        False,
                        "YYYY-MM",
                        "Explicit source manifest competency.",
                    ),
                    field("cod_marca", "Cod_Marca", "integer", False, None, "Brand code."),
                    field(
                        "descr_marca",
                        "Descr_Marca",
                        "string",
                        False,
                        None,
                        "Brand description.",
                    ),
                    field("cod_cargo", "Cod_Cargo", "integer", False, None, "Role code."),
                    field(
                        "descr_cargo",
                        "Descri_Cargo",
                        "string",
                        False,
                        None,
                        "Role description.",
                    ),
                    field(
                        "percentual_comissao",
                        "%_Comiss",
                        "number",
                        False,
                        None,
                        "Fra\u00e7\u00e3o decimal; 0.025 representa 2,5%.",
                    ),
                ],
            },
        },
    }


def field(
    name: str,
    source: str | None,
    type_name: str,
    nullable: bool,
    format_name: str | None,
    semantics: str,
) -> JsonObject:
    result: JsonObject = {
        "name": name,
        "source": source,
        "type": type_name,
        "nullable": nullable,
        "semantics": semantics,
    }
    if format_name is not None:
        result["format"] = format_name
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the canonical Dom Rock dataset.")
    parser.add_argument("--source-root", type=Path, default=DEFAULT_SOURCE_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = build_dataset(source_root=args.source_root, output_dir=args.output_dir)
    print(json.dumps(result.hashes, sort_keys=True))


if __name__ == "__main__":
    main()
