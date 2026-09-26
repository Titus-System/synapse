from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from scripts import build_canonical_dataset as builder
from scripts.build_canonical_dataset import (
    COMISSOES_FIELD_ORDER,
    DEFAULT_OUTPUT_DIR,
    DEFAULT_SOURCE_ROOT,
    PUBLISHED_COMPETENCIAS,
    RH_FIELD_ORDER,
    VENDAS_FIELD_ORDER,
    CompetenciaSource,
    build_dataset,
    build_rh_rows,
    empty_report,
    manifest,
)

COMPETENCIA_PATTERN = re.compile(r"^2025-(08|09|10|11|12)$")
ISO_DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ARTIFACT_NAMES = (
    "rh.jsonl",
    "vendas.jsonl",
    "comissoes.jsonl",
    "schema.json",
    "normalization_report.json",
)


@pytest.fixture(scope="session")
def canonical_dir(tmp_path_factory: pytest.TempPathFactory) -> Path:
    generated_dir = tmp_path_factory.mktemp("canonical")
    build_dataset(source_root=DEFAULT_SOURCE_ROOT, output_dir=generated_dir)

    return generated_dir


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


@pytest.fixture(scope="session")
def rh_rows(canonical_dir: Path) -> list[dict[str, Any]]:
    return read_jsonl(canonical_dir / "rh.jsonl")


@pytest.fixture(scope="session")
def vendas_rows(canonical_dir: Path) -> list[dict[str, Any]]:
    return read_jsonl(canonical_dir / "vendas.jsonl")


@pytest.fixture(scope="session")
def comissoes_rows(canonical_dir: Path) -> list[dict[str, Any]]:
    return read_jsonl(canonical_dir / "comissoes.jsonl")


def test_all_five_artifacts_exist(canonical_dir: Path) -> None:
    assert {path.name for path in canonical_dir.iterdir()} == set(ARTIFACT_NAMES)


def test_versioned_artifacts_match_generated_artifacts(
    tmp_path: Path,
) -> None:
    generated_dir = tmp_path / "generated"

    build_dataset(source_root=DEFAULT_SOURCE_ROOT, output_dir=generated_dir)

    for artifact_name in ARTIFACT_NAMES:
        assert (generated_dir / artifact_name).read_bytes() == (
            DEFAULT_OUTPUT_DIR / artifact_name
        ).read_bytes()


@pytest.mark.parametrize("fixture_name", ["rh_rows", "vendas_rows", "comissoes_rows"])
def test_every_row_has_valid_competencia(request: pytest.FixtureRequest, fixture_name: str) -> None:
    rows = request.getfixturevalue(fixture_name)

    assert all(COMPETENCIA_PATTERN.match(row["competencia"]) for row in rows)


def test_canonical_dates_are_not_excel_serials(
    rh_rows: list[dict[str, Any]],
    vendas_rows: list[dict[str, Any]],
) -> None:
    date_values: list[str] = []
    date_values.extend(row["data_ref"] for row in rh_rows)
    date_values.extend(row["data_admiss"] for row in rh_rows)
    date_values.extend(row["data_demiss"] for row in rh_rows if row["data_demiss"] is not None)
    date_values.extend(row["data_ref"] for row in vendas_rows)
    date_values.extend(row["data_venda"] for row in vendas_rows if row["data_venda"] is not None)

    assert all(isinstance(value, str) for value in date_values)
    assert not any(isinstance(value, int | float) for value in date_values)


def test_canonical_dates_use_iso_format(
    rh_rows: list[dict[str, Any]],
    vendas_rows: list[dict[str, Any]],
) -> None:
    date_values: list[str] = []
    date_values.extend(row["data_ref"] for row in rh_rows)
    date_values.extend(row["data_admiss"] for row in rh_rows)
    date_values.extend(row["data_demiss"] for row in rh_rows if row["data_demiss"] is not None)
    date_values.extend(row["data_ref"] for row in vendas_rows)
    date_values.extend(row["data_venda"] for row in vendas_rows if row["data_venda"] is not None)

    assert all(ISO_DATE_PATTERN.match(value) for value in date_values)


def test_matricula_is_string_and_preserves_matric_prefix(
    rh_rows: list[dict[str, Any]],
    vendas_rows: list[dict[str, Any]],
) -> None:
    matriculas = [row["matricula"] for row in rh_rows + vendas_rows]

    assert all(isinstance(matricula, str) for matricula in matriculas)
    assert all(matricula.startswith("MATRIC-") for matricula in matriculas)


def test_percentual_comissao_remains_fraction(comissoes_rows: list[dict[str, Any]]) -> None:
    percentuais = [row["percentual_comissao"] for row in comissoes_rows]

    assert max(percentuais) < 1
    assert Decimal(str(percentuais[0])) == Decimal("0.025")


def test_novembro_reconciles_discarded_count(vendas_rows: list[dict[str, Any]]) -> None:
    assert source_vendas_row_count("2025-11") == 5000
    assert sum(r["competencia"] == "2025-11" for r in vendas_rows) == 4793


def test_novembro_reconciles_discarded_total(vendas_rows: list[dict[str, Any]]) -> None:
    canonical = sum(
        (Decimal(str(r["vlr_venda"])) for r in vendas_rows if r["competencia"] == "2025-11"),
        Decimal(0),
    )
    assert source_vendas_total("2025-11") - Decimal("581519.4559") == canonical


def test_vendas_from_november_24_to_28_have_data_venda(
    vendas_rows: list[dict[str, Any]],
) -> None:
    rows = [
        row
        for row in vendas_rows
        if row["competencia"] == "2025-11" and row["data_ref"] in november_real_sale_dates()
    ]

    assert len(rows) == 679
    assert all(row["data_venda"] == row["data_ref"] for row in rows)


def test_vendas_from_november_first_have_null_data_venda(
    vendas_rows: list[dict[str, Any]],
) -> None:
    rows = [
        row
        for row in vendas_rows
        if row["competencia"] == "2025-11" and row["data_ref"] == "2025-11-01"
    ]

    assert len(rows) == 4114
    assert all(row["data_venda"] is None for row in rows)


def test_other_published_competencias_do_not_have_unexpected_data_venda(
    vendas_rows: list[dict[str, Any]],
) -> None:
    rows = [row for row in vendas_rows if row["competencia"] != "2025-11"]

    assert all(row["data_venda"] is None for row in rows)


def test_cargo_150_preserves_rh_descriptions_and_consolidates_rates(
    rh_rows: list[dict[str, Any]],
    comissoes_rows: list[dict[str, Any]],
) -> None:
    expected = {"GERENTE DE LOJA", "GERENTE QUIOSQUE"}
    rh_descriptions = {row["descr_cargo"] for row in rh_rows if row["cod_cargo"] == 150}
    comissao_descriptions = {
        row["descr_cargo"] for row in comissoes_rows if row["cod_cargo"] == 150
    }

    assert expected <= rh_descriptions
    assert comissao_descriptions == {"GERENTE DE LOJA"}


def test_vendas_reais_column_is_not_in_canonical_rh(rh_rows: list[dict[str, Any]]) -> None:
    assert all("VENDAS R$" not in row for row in rh_rows)


def test_setembro_matric_246_duplicate_is_handled(
    canonical_dir: Path,
    rh_rows: list[dict[str, Any]],
) -> None:
    rows = [
        row
        for row in rh_rows
        if row["competencia"] == "2025-09" and row["matricula"] == "MATRIC-246"
    ]
    report = json.loads((canonical_dir / "normalization_report.json").read_text(encoding="utf-8"))

    assert len(rows) == 1
    discarded = next(
        row for row in report["discarded_rows"] if row.get("matricula") == "MATRIC-246"
    )
    assert discarded["dataset"] == "rh"
    assert discarded["competencia"] == "2025-09"
    assert discarded["source_file"] == "dataset_domrock/BASE RH/BASE RH_SET25.xlsx"
    assert discarded["source_sheet"] == "BASE HC"
    assert discarded["matricula"] == "MATRIC-246"
    assert discarded["reason"] == "duplicate_integral_rh_row"
    assert discarded["source_row"] > discarded["kept_source_row"]


@pytest.mark.parametrize("cargo_code", [100, 150, 200, 300])
def test_data_demiss_rejects_known_cargo_codes_before_date_conversion(
    tmp_path: Path,
    cargo_code: int,
) -> None:
    source_file = tmp_path / f"rh_data_demiss_{cargo_code}.xlsx"
    write_rh_workbook_with_data_demiss(source_file, cargo_code)
    source = CompetenciaSource(
        competencia="2025-07",
        rh_file=source_file,
        rh_sheet="BASE HC",
        vendas_file=tmp_path / "unused.xlsx",
        vendas_sheet="VENDAS",
    )

    with pytest.raises(ValueError) as exc_info:
        build_rh_rows(source, empty_report())

    message = str(exc_info.value)
    assert "Possible Data_Demiss/Cod_Cargo misalignment" in message
    assert "file=" in message
    assert source_file.name in message
    assert "competencia=2025-07" in message
    assert "row=2" in message
    assert f"Data_Demiss={cargo_code}" in message


def test_vendas_duplicates_are_not_removed(vendas_rows: list[dict[str, Any]]) -> None:
    counts = Counter(
        tuple(row[field] for field in VENDAS_FIELD_ORDER if field != "data_venda")
        for row in vendas_rows
    )

    assert any(count > 1 for count in counts.values())
    source_count = sum(
        source_vendas_row_count(competencia) for competencia in PUBLISHED_COMPETENCIAS
    )

    assert len(vendas_rows) == source_count - 664


def test_schema_matches_physical_jsonl_fields(
    canonical_dir: Path,
    rh_rows: list[dict[str, Any]],
    vendas_rows: list[dict[str, Any]],
    comissoes_rows: list[dict[str, Any]],
) -> None:
    schema = json.loads((canonical_dir / "schema.json").read_text(encoding="utf-8"))

    assert [field["name"] for field in schema["tables"]["rh"]["fields"]] == list(rh_rows[0])
    assert [field["name"] for field in schema["tables"]["vendas"]["fields"]] == list(vendas_rows[0])
    assert [field["name"] for field in schema["tables"]["comissoes"]["fields"]] == list(
        comissoes_rows[0]
    )
    assert list(rh_rows[0]) == list(RH_FIELD_ORDER)
    assert list(vendas_rows[0]) == list(VENDAS_FIELD_ORDER)
    assert list(comissoes_rows[0]) == list(COMISSOES_FIELD_ORDER)


def test_comissao_is_expanded_to_120_records(comissoes_rows: list[dict[str, Any]]) -> None:
    assert len(comissoes_rows) == 120


def test_two_independent_runs_generate_identical_artifacts(tmp_path: Path) -> None:
    first_dir = tmp_path / "run_a"
    second_dir = tmp_path / "run_b"

    first = build_dataset(source_root=DEFAULT_SOURCE_ROOT, output_dir=first_dir)
    second = build_dataset(source_root=DEFAULT_SOURCE_ROOT, output_dir=second_dir)

    assert first.hashes == second.hashes
    for artifact_name in ARTIFACT_NAMES:
        assert (first_dir / artifact_name).read_bytes() == (second_dir / artifact_name).read_bytes()


def test_report_includes_data_venda_zero_counts_for_all_competencias(
    canonical_dir: Path,
) -> None:
    report = json.loads((canonical_dir / "normalization_report.json").read_text(encoding="utf-8"))

    assert report["checks"]["vendas"]["data_venda_rows_by_competencia"] == {
        "2025-08": 0,
        "2025-09": 0,
        "2025-10": 0,
        "2025-11": 679,
        "2025-12": 0,
    }


def source_vendas_row_count(competencia: str) -> int:
    source = next(item for item in manifest(DEFAULT_SOURCE_ROOT) if item.competencia == competencia)
    return len(read_source_rows(source.vendas_file, source.vendas_sheet))


def source_vendas_total(competencia: str) -> Decimal:
    source = next(item for item in manifest(DEFAULT_SOURCE_ROOT) if item.competencia == competencia)
    return sum(
        (
            Decimal(str(row["Vlr _Venda"]))
            for row in read_source_rows(source.vendas_file, source.vendas_sheet)
        ),
        Decimal("0"),
    )


def read_source_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(rows)]
        return [
            dict(zip(headers, row, strict=True))
            for row in rows
            if not all(value is None for value in row)
        ]
    finally:
        workbook.close()


def write_rh_workbook_with_data_demiss(path: Path, data_demiss: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "BASE HC"
    worksheet.append(list(RH_SOURCE_HEADERS))
    worksheet.append(
        [
            date(2025, 7, 1),
            1,
            "MARCA",
            101,
            "LOJA",
            "MATRIC-001",
            date(2025, 1, 1),
            data_demiss,
            100,
            "VENDEDOR",
        ]
    )
    workbook.save(path)


RH_SOURCE_HEADERS = (
    "Data_Ref",
    "Cod_Marca",
    "Descri_Marca",
    "Cod_Loja",
    "Descr_Loja",
    "Matricula",
    "Data_Admiss",
    "Data_Demiss",
    "Cod_Cargo",
    "Descri_Cargo",
)


@pytest.fixture(scope="session")
def report(canonical_dir: Path) -> dict[str, Any]:
    return cast(
        dict[str, Any],
        json.loads((canonical_dir / "normalization_report.json").read_text(encoding="utf-8")),
    )


@pytest.fixture(scope="session")
def source_rh() -> list[builder.CanonicalRow]:
    diagnostic = builder.empty_report()
    return [
        row for source in manifest(DEFAULT_SOURCE_ROOT) for row in build_rh_rows(source, diagnostic)
    ]


def rh_record(
    competencia: str,
    matricula: str = "TEST-EMPLOYEE",
    source_row: int = 2,
    **changes: builder.JsonValue,
) -> builder.CanonicalRow:
    record: builder.JsonObject = {
        "competencia": competencia,
        "data_ref": competencia + "-01",
        "cod_marca": 10,
        "descr_marca": "PRETO",
        "cod_loja": 1,
        "descr_loja": "LOJA-1",
        "matricula": matricula,
        "data_admiss": "2020-01-01",
        "data_demiss": None,
        "cod_cargo": 100,
        "descr_cargo": "VENDEDOR LOJA",
        **changes,
    }
    return builder.CanonicalRow(
        DEFAULT_SOURCE_ROOT / f"rh-{competencia}.xlsx", "BASE HC", source_row, record
    )


@pytest.mark.parametrize("fixture_name", ["rh_rows", "vendas_rows", "comissoes_rows"])
def test_exactly_five_published_competencias(
    request: pytest.FixtureRequest,
    fixture_name: str,
) -> None:
    rows = request.getfixturevalue(fixture_name)
    assert {row["competencia"] for row in rows} == set(PUBLISHED_COMPETENCIAS)


def test_observed_synthetic_and_final_counts(
    rh_rows: list[dict[str, Any]], vendas_rows: list[dict[str, Any]], report: dict[str, Any]
) -> None:
    population = report["rh_population"]
    assert population["observed_row_count"] == 2328
    assert population["synthetic_row_count"] == 418
    assert population["canonical_row_count"] == len(rh_rows) == 2746
    assert population["fills_by_direction"] == {"previous": 384, "next": 34}
    assert Counter(r["competencia"] for r in rh_rows) == dict(
        zip(PUBLISHED_COMPETENCIAS, [500, 539, 551, 568, 588], strict=True)
    )
    assert len(vendas_rows) == 24336
    assert report["checks"]["rh"]["total_rows"] == 2746


def test_july_is_excluded_but_still_analyzed(report: dict[str, Any]) -> None:
    july = report["competency_status"]["2025-07"]
    assert july["status"] == "excluded"
    assert july["published"] is False
    assert july["historical_evidence"] is True
    assert july["excluded_rh_rows"] == 322
    assert july["excluded_vendas_rows"] == 5018
    assert len(july["august_matriculas_missing_in_july"]) == 167
    assert len(july["admitted_before_july_missing_in_july"]) == 141
    assert report["source_validations"]["rh"]["2025-07"]["canonical_row_count"] == 0
    assert report["source_validations"]["vendas"]["2025-07"]["canonical_total_vlr_venda"] == "0"
    decisions = {d["id"]: d for d in report["decisions"]}
    assert decisions["EXCLUDE_2025_07"]["status"] == "defined"
    assert decisions["EXCLUDE_2025_07"]["reason"] == "unreconcilable_source_population"


def test_generic_rh_uniqueness(rh_rows: list[dict[str, Any]]) -> None:
    assert len({(r["competencia"], r["matricula"]) for r in rh_rows}) == len(rh_rows)


def test_generic_integral_duplicates_preserve_first_and_report_every_excess() -> None:
    first = rh_record("2025-10")
    second = rh_record("2025-10", source_row=6)
    third = rh_record("2025-10", source_row=9)
    other_month = rh_record("2025-11")
    diagnostic = builder.empty_report()
    rows = builder.deduplicate_rh([first, second, other_month, third], diagnostic)
    assert rows == [first, other_month]
    discarded = builder.discarded_rows(diagnostic)
    assert [r["source_row"] for r in discarded] == [6, 9]
    assert all(r["kept_source_row"] == 2 for r in discarded)
    assert all(r["source_sheet"] == "BASE HC" for r in discarded)
    assert all(r["matricula"] == "TEST-EMPLOYEE" for r in discarded)


def test_conflicting_rh_duplicates_fail_with_identity_and_source_rows() -> None:
    rows = [rh_record("2025-10"), rh_record("2025-10", source_row=7, cod_loja=2)]
    with pytest.raises(ValueError, match="Conflicting RH rows") as error:
        builder.deduplicate_rh(rows, builder.empty_report())
    message = str(error.value)
    assert all(value in message for value in ["2025-10", "TEST-EMPLOYEE", "2", "7", "BASE HC"])


@pytest.mark.parametrize("dataset", ["rh", "vendas"])
@pytest.mark.parametrize("invalid_date", [date(2025, 10, 1), date(2024, 11, 1)])
def test_build_rejects_reference_month_or_year_mismatch(
    dataset: str,
    invalid_date: date,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = next(s for s in manifest(DEFAULT_SOURCE_ROOT) if s.competencia == "2025-11")
    path = source.rh_file if dataset == "rh" else source.vendas_file
    sheet = source.rh_sheet if dataset == "rh" else source.vendas_sheet
    table = builder.read_sheet(path, sheet)
    values = table.rows[0].values.copy()
    values["Data_Ref" if dataset == "rh" else "Date_Ref"] = invalid_date
    invalid = builder.SourceTable(table.headers, (builder.SourceRow(2, values),))
    monkeypatch.setattr(builder, "read_sheet", lambda *_: invalid)
    build_rows = builder.build_rh_rows if dataset == "rh" else builder.build_vendas_rows
    with pytest.raises(ValueError, match="data_ref month differs from manifest") as error:
        build_rows(source, builder.empty_report())
    assert "competencia=2025-11" in str(error.value)
    assert "row=2" in str(error.value)


def test_admissions_are_global_minimum_and_other_attributes_unchanged(
    source_rh: list[builder.CanonicalRow],
    rh_rows: list[dict[str, Any]],
    report: dict[str, Any],
) -> None:
    history = builder.build_history(source_rh)
    originals = {(r.record["competencia"], r.record["matricula"]): r.record for r in source_rh}
    for row in rh_rows:
        assert row["data_admiss"] == min(
            cast(str, r.record["data_admiss"]) for r in history[row["matricula"]]
        )
        original = originals.get((row["competencia"], row["matricula"]))
        if original is None:
            continue
        for field in builder.RH_FIELD_ORDER:
            if field not in {"data_admiss", "data_demiss"}:
                assert row[field] == original[field]
    assert report["source_validations"]["history"]["matriculas_with_admission_variations"] == 27
    reconciliations = [
        r for r in report["reconciliations"] if r["type"] == "canonical_admission_date"
    ]
    assert len(reconciliations) == 66
    assert all(
        {
            "evidence_competencia",
            "evidence_source_file",
            "evidence_source_sheet",
            "evidence_source_row",
            "original_value",
            "canonical_value",
        }
        <= r.keys()
        for r in reconciliations
    )
    assert all(r["original_value"] != r["canonical_value"] for r in reconciliations)


@pytest.mark.parametrize(
    ("matricula", "expected"),
    [
        ("MATRIC-551", "2002-04-27"),
        ("MATRIC-257", "2025-06-06"),
        ("MATRIC-50", "2024-07-01"),
    ],
)
def test_admission_can_come_from_july_or_later_month(
    matricula: str,
    expected: str,
    rh_rows: list[dict[str, Any]],
) -> None:
    rows = [r for r in rh_rows if r["matricula"] == matricula]
    assert rows
    assert {r["data_admiss"] for r in rows} == {expected}


def test_terminal_termination_regression(
    source_rh: list[builder.CanonicalRow],
    rh_rows: list[dict[str, Any]],
    report: dict[str, Any],
) -> None:
    history = builder.build_history(source_rh)
    for row in rh_rows:
        terminations = [
            cast(str, r.record["data_demiss"])
            for r in history[row["matricula"]]
            if r.record["data_demiss"] is not None
        ]
        if terminations:
            earliest = min(terminations)
            if row["competencia"] >= earliest[:7]:
                assert row["data_demiss"] == earliest
    changes = [
        r for r in report["reconciliations"] if r["type"] == "terminal_termination_propagation"
    ]
    assert len(changes) == 71
    assert len({r["matricula"] for r in changes}) == 26
    assert sum(r["original_value"] is None for r in changes) == 70
    second_termination = next(r for r in changes if r["original_value"] == "2025-12-05")
    assert second_termination["matricula"] == "MATRIC-241"
    assert second_termination["canonical_value"] == "2025-08-31"
    assert second_termination["origin_competencia"] == "2025-08"
    assert sum(r["origin_competencia"] == "2025-07" for r in changes) == 16


def test_temporal_reconciliation_handles_gaps_and_preserves_source() -> None:
    source = [
        rh_record("2025-07", data_demiss="2025-07-15"),
        rh_record("2025-09", cod_loja=7),
        rh_record("2025-12", data_admiss="1999-01-01", data_demiss="2025-12-05"),
    ]
    diagnostic = builder.empty_report()
    result = builder.reconcile_rh(builder.build_history(source), source, diagnostic)
    assert len(result) == 2
    assert {r.record["data_admiss"] for r in result} == {"1999-01-01"}
    assert {r.record["data_demiss"] for r in result} == {"2025-07-15"}
    assert result[0].record["cod_loja"] == 7
    assert source[1].record["data_demiss"] is None
    assert source[1].record["data_admiss"] == "2020-01-01"


def test_termination_does_not_change_competencies_before_termination() -> None:
    source = [
        rh_record("2025-08"),
        rh_record("2025-10", data_demiss="2025-10-05"),
        rh_record("2025-12"),
    ]
    result = builder.reconcile_rh(builder.build_history(source), source, builder.empty_report())
    assert [r.record["data_demiss"] for r in result] == [None, "2025-10-05", "2025-10-05"]


def test_post_termination_sales_discarded_and_auditable(report: dict[str, Any]) -> None:
    rows = [r for r in report["discarded_rows"] if r["reason"] == "sale_after_terminal_termination"]
    assert len(rows) == 656
    assert sum((Decimal(str(r["vlr_venda"])) for r in rows), Decimal(0)) == Decimal(
        "1667160.9980997"
    )
    assert all(
        {
            "source_file",
            "source_sheet",
            "source_row",
            "competencia",
            "matricula",
            "data_ref",
            "data_venda",
            "cod_marca",
            "cod_loja",
            "vlr_venda",
            "terminal_termination_date",
        }
        <= r.keys()
        for r in rows
    )
    assert report["sale_discard_summary"]["dated_same_termination_month_discarded"] == 8
    assert report["sale_discard_summary"]["undated_same_termination_month_preserved"] == 229


def test_commission_unique_keys_matches_and_canonical_manager_rates(
    rh_rows: list[dict[str, Any]],
    comissoes_rows: list[dict[str, Any]],
    report: dict[str, Any],
) -> None:
    rates = {10: 0.01, 20: 0.015, 30: 0.005, 40: 0.02, 50: 0.01, 60: 0.015}
    index = Counter(tuple(row[field] for field in builder.COMISSAO_KEY) for row in comissoes_rows)
    assert len(index) == 120
    assert set(index.values()) == {1}
    assert Counter(row["competencia"] for row in comissoes_rows) == dict.fromkeys(
        PUBLISHED_COMPETENCIAS, 24
    )
    for row in rh_rows:
        assert index[tuple(row[field] for field in builder.COMISSAO_KEY)] == 1
    for row in comissoes_rows:
        if row["cod_cargo"] == 150:
            assert row["descr_cargo"] == "GERENTE DE LOJA"
            assert row["percentual_comissao"] == rates[row["cod_marca"]]
    discarded = [r for r in report["discarded_rows"] if r["dataset"] == "comissoes"]
    assert len(discarded) == 6
    assert {r["cod_marca"] for r in discarded} == set(rates)
    assert all(r["descr_cargo"] == "GERENTE QUIOSQUE" and r["source_row"] >= 2 for r in discarded)
    assert all(
        r["reason"] == "manager_kiosk_rate_superseded_by_canonical_manager_rate" for r in discarded
    )


def test_commission_duplicate_or_missing_match_fails() -> None:
    row = rh_record("2025-08")
    with pytest.raises(ValueError, match="Commission keys must be unique"):
        builder.validate_commission_matches([], [row, row])
    with pytest.raises(ValueError, match="Expected exactly one commission"):
        builder.validate_commission_matches([row], [])


@pytest.mark.parametrize("competencia", PUBLISHED_COMPETENCIAS)
def test_retained_sales_match_source_and_discard_ledger_exactly(
    competencia: str, vendas_rows: list[dict[str, Any]], report: dict[str, Any]
) -> None:
    source = next(s for s in manifest(DEFAULT_SOURCE_ROOT) if s.competencia == competencia)
    table = builder.read_sheet(source.vendas_file, source.vendas_sheet)
    discarded = [
        r
        for r in report["discarded_rows"]
        if r["dataset"] == "vendas" and r["competencia"] == competencia
    ]
    excluded = {r["source_row"] for r in discarded}
    assert len(excluded) == len(discarded)
    retained = [
        builder.build_venda_row(source, row).record
        for row in table.rows
        if row.row_number not in excluded
    ]
    published = [r for r in vendas_rows if r["competencia"] == competencia]
    assert published == retained
    assert len(published) + len(discarded) == 5000
    total = sum((Decimal(str(r["vlr_venda"])) for r in published + discarded), Decimal(0))
    assert total == source_vendas_total(competencia)


def test_orphans_reconstructed_or_discarded(
    report: dict[str, Any], rh_rows: list[dict[str, Any]], vendas_rows: list[dict[str, Any]]
) -> None:
    orphan = next(i for i in report["invariants"] if i["id"] == "I2")
    assert orphan["violations_found"] == orphan["reconciliations_applied"] == 60
    assert orphan["remaining_violations"] == 0 and not orphan["blocking"]
    assert Counter(r["action"] for r in orphan["details"]) == {
        "reconstructed_rh": 52,
        "sale_without_reconstructable_rh": 8,
    }
    discarded = [
        r for r in report["discarded_rows"] if r["reason"] == "sale_without_reconstructable_rh"
    ]
    assert Counter(r["matricula"] for r in discarded) == {
        "MATRIC-599": 6,
        "MATRIC-191": 1,
        "MATRIC-207": 1,
    }
    keys = Counter((r["competencia"], r["matricula"]) for r in rh_rows)
    assert all(keys[r["competencia"], r["matricula"]] == 1 for r in vendas_rows)


def test_sale_dimensions_are_authoritative_nonblocking(report: dict[str, Any]) -> None:
    dimensions = next(i for i in report["invariants"] if i["id"] == "I3")
    assert not dimensions["blocking"]
    assert dimensions["status"] == "passed"
    assert dimensions["remaining_violations"] == 0
    source = report["source_validations"]["sale_dimensions"]
    assert sum(d["sale_rows"] for d in source["by_competencia"]) == 241
    assert (
        sum(
            d["sale_rows"]
            for d in source["by_competencia"]
            if d["competencia"] in PUBLISHED_COMPETENCIAS
        )
        == 77
    )
    assert sum(d["sale_rows"] for d in source["final_by_competencia"]) == 74
    assert any(d["multiple_store_sales"] for d in dimensions["details"])


def test_attribute_changes_are_observations_with_no_fabricated_events(
    report: dict[str, Any],
) -> None:
    changes = [
        w for w in report["warnings"] if w.get("classification") == "observed_attribute_change"
    ]
    assert Counter(w["field"] for w in changes) == {
        "cod_loja": 38,
        "cod_marca": 6,
        "cod_cargo": 11,
        "descr_cargo": 16,
    }
    assert all(w["action"] == "preserved" and w["effective_date"] is None for w in changes)
    assert all(w["blocking"] is False for w in changes)
    gap = next(
        w
        for w in changes
        if w["matricula"] == "MATRIC-429"
        and w["field"] == "cod_loja"
        and w["current_competencia"] == "2025-12"
    )
    assert gap["intermediate_missing_competencias"] == ["2025-09", "2025-10", "2025-11"]


def test_rh_continuity_covers_all_eligible_slots(
    report: dict[str, Any], source_rh: list[builder.CanonicalRow], rh_rows: list[dict[str, Any]]
) -> None:
    invariant = next(i for i in report["invariants"] if i["id"] == "I7")
    assert not invariant["blocking"]
    assert invariant["violations_found"] == invariant["reconciliations_applied"] == 418
    assert invariant["remaining_violations"] == 0
    keys = {(r["competencia"], r["matricula"]) for r in rh_rows}
    originals = {(r.record["competencia"], r.record["matricula"]) for r in source_rh}
    history = builder.build_history(source_rh)
    for c in PUBLISHED_COMPETENCIAS:
        for m, records in history.items():
            admission = min(str(r.record["data_admiss"]) for r in records)
            dates = [str(r.record["data_demiss"]) for r in records if r.record["data_demiss"]]
            eligible = admission[:7] <= c and (not dates or min(dates)[:7] >= c)
            assert ((c, m) in keys) == (eligible or (c, m) in originals)
    index = {(r["competencia"], r["matricula"]): r for r in rh_rows}
    assert index["2025-08", "MATRIC-365"]["cod_cargo"] == 200
    assert index["2025-12", "MATRIC-478"]["cod_loja"] == 38


def test_report_states_and_required_structure(report: dict[str, Any]) -> None:
    assert report["version"] == 3
    assert {
        "source_validations",
        "normalizations_applied",
        "discarded_rows",
        "discarded_columns",
        "reconciliations",
        "warnings",
        "invariants",
        "competency_status",
        "decisions",
        "checks",
    } <= report.keys()
    invariants = {i["id"]: i for i in report["invariants"]}
    assert list(invariants) == [f"I{n}" for n in range(1, 10)]
    assert {key: value["status"] for key, value in invariants.items()} == {
        "I1": "reconciled",
        "I2": "reconciled",
        "I3": "passed",
        "I4": "reconciled",
        "I5": "warning",
        "I6": "reconciled",
        "I7": "reconciled",
        "I8": "passed",
        "I9": "passed",
    }
    for item in invariants.values():
        assert {
            "name",
            "blocking",
            "count_unit",
            "violations_found",
            "reconciliations_applied",
            "remaining_violations",
            "details",
        } <= item.keys()
        assert item["violations_found"] == (
            item["reconciliations_applied"] + item["remaining_violations"]
        )
    for month in PUBLISHED_COMPETENCIAS:
        status = report["competency_status"][month]
        assert status["status"] == "ready" and status["published"] is True
        assert status["blocking_invariants"] == []
        assert "I4" not in status["blocking_invariants"]
        assert status["blocking_decisions"] == []
    serialized = json.dumps(report)
    assert str(DEFAULT_SOURCE_ROOT) not in serialized
    assert "C:\\\\" not in serialized
    assert all(d["status"] == "defined" for d in report["decisions"])


def test_schema_declares_temporal_and_commission_semantics(canonical_dir: Path) -> None:
    schema = json.loads((canonical_dir / "schema.json").read_text(encoding="utf-8"))
    assert schema["version"] == 3
    assert schema["published_competencias"] == list(PUBLISHED_COMPETENCIAS)
    assert schema["historical_evidence"]["excluded_competencias"] == ["2025-07"]
    assert schema["tables"]["rh"]["unique_by"] == ["competencia", "matricula"]
    commission = schema["tables"]["comissoes"]
    assert commission["unique_by"] == list(builder.COMISSAO_KEY)
    assert commission["manager_rate"]["source_description"] == "GERENTE DE LOJA"
    assert schema["tables"]["rh"]["commission_join"]["descr_cargo_participates"] is False
    fields = {f["name"]: f for f in schema["tables"]["rh"]["fields"]}
    assert "ineligible" in fields["data_demiss"]["semantics"]
    assert "including July" in fields["data_admiss"]["semantics"]


def november_real_sale_dates() -> set[str]:
    return {
        "2025-11-24",
        "2025-11-25",
        "2025-11-26",
        "2025-11-27",
        "2025-11-28",
    }


@pytest.mark.parametrize(
    "evidence_month, admission, expected",
    [
        ("2025-07", "2020-01-01", ["2025-08", "2025-09", "2025-10", "2025-11", "2025-12"]),
        ("2025-10", "2025-09-30", ["2025-09", "2025-10", "2025-11", "2025-12"]),
    ],
)
def test_continuity_uses_only_observed_evidence(
    evidence_month: str, admission: str, expected: list[str]
) -> None:
    source = [rh_record(evidence_month, data_admiss=admission)]
    history = builder.build_history(source)
    diagnostic = builder.empty_report()
    observed = builder.reconcile_rh(history, source, diagnostic)
    result = builder.fill_rh_continuity(history, observed, diagnostic)
    assert [r.record["competencia"] for r in result] == expected
    fills = builder.reconciliations_of(diagnostic, "rh_continuity_fill")
    assert all(r["evidence_competencia"] == evidence_month for r in fills)
    assert all(
        r["source_direction"]
        == ("previous" if cast(str, r["competencia"]) > evidence_month else "next")
        for r in fills
    )
    assert all(r["reason"] == "maintain_employee_continuity" for r in fills)
    assert all(r.record["data_ref"] == str(r.record["competencia"]) + "-01" for r in result)


def test_gap_prefers_latest_observed_previous_attributes_over_future() -> None:
    source = [
        rh_record("2025-07", cod_cargo=200, cod_loja=3),
        rh_record("2025-10", cod_cargo=100, cod_loja=9),
    ]
    history = builder.build_history(source)
    report = builder.empty_report()
    rows = builder.fill_rh_continuity(
        history, builder.reconcile_rh(history, source, report), report
    )
    assert [(r.record["cod_cargo"], r.record["cod_loja"]) for r in rows] == [
        (200, 3),
        (200, 3),
        (100, 9),
        (100, 9),
        (100, 9),
    ]
    fills = builder.reconciliations_of(report, "rh_continuity_fill")
    assert [r["evidence_competencia"] for r in fills] == [
        "2025-07",
        "2025-07",
        "2025-10",
        "2025-10",
    ]


def test_no_synthetic_rh_after_termination_but_observed_rows_remain() -> None:
    source = [rh_record("2025-07", data_demiss="2025-09-15"), rh_record("2025-12")]
    report = builder.empty_report()
    history = builder.build_history(source)
    rows = builder.fill_rh_continuity(
        history, builder.reconcile_rh(history, source, report), report
    )
    assert [(r.record["competencia"], r.record["data_demiss"]) for r in rows] == [
        ("2025-08", None),
        ("2025-09", "2025-09-15"),
        ("2025-12", "2025-09-15"),
    ]
    assert source[1].record["data_demiss"] is None


@pytest.mark.parametrize(
    "competencia, actual_date, expected_discard",
    [
        ("2025-11", "2025-11-24", True),
        ("2025-11", "2025-11-20", False),
        ("2025-11", "2025-11-19", False),
        ("2025-11", None, False),
        ("2025-12", None, True),
        ("2025-10", None, False),
    ],
)
def test_sale_termination_uses_actual_date_or_month_only(
    competencia: str, actual_date: str | None, expected_discard: bool
) -> None:
    source = [rh_record("2025-11", data_demiss="2025-11-20")]
    sale = rh_record(competencia, data_venda=actual_date, vlr_venda=0.1, cod_loja=99, cod_marca=60)
    report = builder.empty_report()
    result = builder.reconcile_sales(builder.build_history(source), [sale], [sale], report)
    assert result == ([] if expected_discard else [sale])
    assert len(builder.discarded_rows(report)) == int(expected_discard)
    if result:
        assert result[0].record["cod_loja"] == 99
        assert result[0].record["cod_marca"] == 60


def test_sale_discard_priority_and_unknown_employee() -> None:
    source = [rh_record("2025-07", data_demiss="2025-07-15")]
    sales = [
        rh_record("2025-08", data_venda=None, vlr_venda=0.1),
        rh_record("2025-08", matricula="NO-HISTORY", source_row=3, data_venda=None, vlr_venda=0.2),
    ]
    report = builder.empty_report()
    assert builder.reconcile_sales(builder.build_history(source), [], sales, report) == []
    assert [r["reason"] for r in builder.discarded_rows(report)] == [
        "sale_after_terminal_termination",
        "sale_without_reconstructable_rh",
    ]


def test_exact_decimal_reconciliation_and_mismatch_rejected() -> None:
    first = rh_record("2025-08", vlr_venda=0.1)
    second = rh_record("2025-08", source_row=3, vlr_venda=0.2)
    detail = builder.sales_reconciliation([first, second], [second], [first.record])
    assert detail["source_total_vlr_venda"] == "0.3"
    assert detail["discarded_total_vlr_venda"] == "0.1"
    assert detail["canonical_total_vlr_venda"] == "0.2"
    assert detail["reconciliation_difference"] == "0"
    with pytest.raises(ValueError, match="minus discarded"):
        builder.sales_reconciliation([first, second], [first], [first.record])


def test_financial_report_reconciles_each_month_and_aggregate(report: dict[str, Any]) -> None:
    financial = report["financial_reconciliation"]
    for detail in [*financial["by_competencia"], financial["aggregate"]]:
        assert (
            detail["source_row_count"] - detail["discarded_row_count"]
            == detail["canonical_row_count"]
        )
        assert Decimal(detail["source_total_vlr_venda"]) - Decimal(
            detail["discarded_total_vlr_venda"]
        ) == Decimal(detail["canonical_total_vlr_venda"])
        assert detail["reconciliation_difference"] == "0"
    assert financial["aggregate"]["canonical_total_vlr_venda"] == "55267944.1007781000001"
